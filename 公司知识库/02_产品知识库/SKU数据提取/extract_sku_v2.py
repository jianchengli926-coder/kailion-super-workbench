#!/usr/bin/env python3
"""SKU结构化提取v2 - 修正列映射
布局: D列(index3)=标签, E列(index4)=值, F列(index5)=标签, G列(index6)=值
       A列(index0)=左侧标签, B/C列(index1/2)=值
"""
import openpyxl
import os
import json
import re

BASE = "/Volumes/Kingston 1TB NV1 40Gbps/豆包独立站SEO项目"
OUT_DIR = os.path.join(BASE, "公司知识库/02_产品知识库/SKU数据提取")

MAIN_XLSX = {
    "剪刀SKU.xlsx": os.path.join(BASE, "剪刀SKU.xlsx"),
    "户外刀（主打刀）SKU.xlsx": os.path.join(BASE, "户外刀（主打刀）SKU.xlsx"),
    "厨刀套装SKU.xlsx": os.path.join(BASE, "厨刀套装SKU.xlsx"),
    "厨刀SKU.xlsx": os.path.join(BASE, "厨刀SKU.xlsx"),
    "户外刀仿牌sku.xlsx": os.path.join(BASE, "户外刀仿牌sku.xlsx"),
    "户外刀Asku.xlsx": os.path.join(BASE, "户外刀Asku.xlsx"),
    "厨房配件SKU.xlsx": os.path.join(BASE, "厨房配件SKU.xlsx"),
}

SENSITIVE_KEYWORDS = ["价格区间", "建议售价", "成本", "出厂价", "进货价", "采购价", "利润", "wholesale price", "cost", "FOB"]

def is_sensitive(label):
    if not label:
        return False
    ll = label.lower()
    for kw in SENSITIVE_KEYWORDS:
        if kw.lower() in ll:
            return True
    return False

def parse_sku_sheet(ws, source_file, sheet_name):
    """解析SKU工作表"""
    record = {
        "sku_sheet_name": sheet_name.strip(),
        "source_file": source_file,
    }
    rows = list(ws.iter_rows(values_only=True))

    for row in rows:
        cells = list(row) + [None] * (10 - len(row))
        # 列映射: A=0,B=1,C=2,D=3,E=4,F=5,G=6,H=7,I=8
        # 主布局: D=label, E=value, F=label, G=value
        d_label = cells[3] if len(cells) > 3 else None
        e_val = cells[4] if len(cells) > 4 else None
        f_label = cells[5] if len(cells) > 5 else None
        g_val = cells[6] if len(cells) > 6 else None

        # D=label, E=value
        if d_label and str(d_label).strip():
            label = str(d_label).strip()
            val = str(e_val).strip() if e_val is not None else ""
            if label and val:
                record[label] = val

        # F=label, G=value
        if f_label and str(f_label).strip():
            label = str(f_label).strip()
            val = str(g_val).strip() if g_val is not None else ""
            if label and val:
                record[label] = val

        # 左侧: A=label, B或C=value
        a_label = cells[0] if len(cells) > 0 else None
        b_val = cells[1] if len(cells) > 1 else None
        c_val = cells[2] if len(cells) > 2 else None

        if a_label and str(a_label).strip():
            label = str(a_label).strip()
            # 跳过section标题
            if label in ["基础SKU", "SEO命名", "市场定位", "中英文描述", "识别特征"]:
                continue
            # 值可能在B或C列
            val = ""
            if b_val and not str(b_val).startswith("=DISPIMG"):
                val = str(b_val).strip()
            elif c_val and not str(c_val).startswith("=DISPIMG"):
                val = str(c_val).strip()
            if label and val and label not in record:
                record[label] = val

    # 清理敏感值
    safe_record = {}
    sensitive_found = []
    for k, v in record.items():
        if k.startswith("_"):
            continue
        if is_sensitive(k):
            sensitive_found.append(k)
            safe_record[k] = "[CONFIDENTIAL-存在]"
        else:
            safe_record[k] = v
    safe_record["_sensitive_fields"] = sensitive_found
    return safe_record

# ========== 主流程 ==========
all_skus = []
non_sku_sheets = []

for fname, fpath in MAIN_XLSX.items():
    if not os.path.exists(fpath):
        continue
    wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
    count = 0
    for sname in wb.sheetnames:
        sname_s = sname.strip()
        if sname_s.startswith("KL-"):
            ws = wb[sname]
            rec = parse_sku_sheet(ws, fname, sname_s)
            all_skus.append(rec)
            count += 1
        else:
            non_sku_sheets.append({"file": fname, "sheet": sname})
    wb.close()
    print(f"{fname}: {count} SKU")

print(f"\n总计: {len(all_skus)} SKU工作表")

# 字段统计
field_stats = {}
for rec in all_skus:
    for k, v in rec.items():
        if k.startswith("_") or k in ["sku_sheet_name", "source_file"]:
            continue
        if v and v != "[CONFIDENTIAL-存在]":
            field_stats[k] = field_stats.get(k, 0) + 1

print("\n=== 字段出现率(有值) ===")
for k, v in sorted(field_stats.items(), key=lambda x: -x[1]):
    pct = round(v / len(all_skus) * 100, 1)
    sens = " [敏感]" if is_sensitive(k) else ""
    print(f"  {k}: {v}/{len(all_skus)} ({pct}%){sens}")

# 保存
with open(os.path.join(OUT_DIR, "_all_sku_records.json"), "w", encoding="utf-8") as f:
    json.dump(all_skus, f, ensure_ascii=False, indent=2)

# SKU列表
sku_list = []
for rec in all_skus:
    sku_code = rec.get("完整SKU", rec["sku_sheet_name"])
    sku_list.append({
        "sku": sku_code,
        "sheet_name": rec["sku_sheet_name"],
        "source_file": rec["source_file"],
        "cn_name": rec.get("中文名称", ""),
        "en_name": rec.get("英文名称", ""),
        "category_cn": rec.get("品类", ""),
        "category_en": rec.get("CATEGORY (EN)", ""),
        "blade_material": rec.get("刀身材质", ""),
        "handle_material": rec.get("柄材", ""),
        "size": rec.get("尺寸", ""),
        "color": rec.get("颜色", ""),
        "finish": rec.get("表面工艺", ""),
        "moq": rec.get("MOQ", ""),
    })

with open(os.path.join(OUT_DIR, "_sku_list.json"), "w", encoding="utf-8") as f:
    json.dump(sku_list, f, ensure_ascii=False, indent=2)

# 抽查5条
print("\n=== 抽查5条 ===")
import random
random.seed(42)
sample = random.sample(all_skus, 5)
for rec in sample:
    sku = rec.get("完整SKU", rec["sku_sheet_name"])
    cn = rec.get("中文名称", "—")
    en = rec.get("英文名称", "—")
    mat = rec.get("刀身材质", "—")
    handle = rec.get("柄材", "—")
    moq = rec.get("MOQ", "—")
    sens = rec.get("_sensitive_fields", [])
    print(f"  {sku}")
    print(f"    CN: {cn[:40]}")
    print(f"    EN: {en[:50]}")
    print(f"    材质: {mat[:30]} | 柄: {handle[:30]} | MOQ: {moq}")
    print(f"    敏感字段: {sens}")
