#!/usr/bin/env python3
"""SKU数据提取脚本 - v7.0 Phase2 Part3
只提取表格中明确存在的数据，不补全、不推测。
"""
import openpyxl
import os
import json
import csv
import re
from pathlib import Path
from collections import defaultdict

BASE = "/Volumes/Kingston 1TB NV1 40Gbps/豆包独立站SEO项目"
OUT_DIR = os.path.join(BASE, "公司知识库/02_产品知识库/SKU数据提取")

# === 主xlsx文件列表 ===
MAIN_XLSX = [
    os.path.join(BASE, "剪刀SKU.xlsx"),
    os.path.join(BASE, "户外刀（主打刀）SKU.xlsx"),
    os.path.join(BASE, "厨刀套装SKU.xlsx"),
    os.path.join(BASE, "厨刀SKU.xlsx"),
    os.path.join(BASE, "户外刀仿牌sku.xlsx"),
    os.path.join(BASE, "户外刀Asku.xlsx"),
    os.path.join(BASE, "厨房配件SKU.xlsx"),
]

# === 汇总/SEO xlsx ===
SUMMARY_XLSX = [
    os.path.join(BASE, "四大品类sku/四大品类_SKU_SEO产品目录集合.xlsx"),
    os.path.join(BASE, "四大品类sku/厨房用品与多功能剪刀_SKU_SEO总表.xlsx"),
    os.path.join(BASE, "四大品类sku/厨房用品与多功能剪刀_SKU_SEO总表_无图扩展版.xlsx"),
    os.path.join(BASE, "四大品类_SKU_SEO产品目录集合_副本.xlsx"),
    os.path.join(BASE, "product-category-mapping.xlsx"),
]

# 敏感字段关键词
SENSITIVE_KEYWORDS = ["成本", "出厂价", "进货价", "采购价", "利润", "cost", "price", "wholesale", "FOB", "报价", "单价", "售价"]

def is_sensitive_col(col_name):
    if not col_name:
        return False
    cn = str(col_name).lower()
    for kw in SENSITIVE_KEYWORDS:
        if kw.lower() in cn:
            return True
    return False

def read_xlsx_info(filepath):
    """读取xlsx文件的工作表信息、表头、行数"""
    result = {
        "filepath": filepath,
        "filename": os.path.basename(filepath),
        "size_mb": round(os.path.getsize(filepath) / 1024 / 1024, 2),
        "sheets": []
    }
    try:
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            rows = list(ws.iter_rows(max_row=3, values_only=True))
            header = list(rows[0]) if rows else []
            # 统计行数（粗略：迭代计数）
            row_count = 0
            for _ in ws.iter_rows(min_row=2, values_only=True):
                row_count += 1
            # 清理表头None
            header_clean = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(header)]
            # 检测敏感列
            sensitive_cols = [h for h in header_clean if is_sensitive_col(h)]
            result["sheets"].append({
                "sheet_name": sname,
                "header": header_clean,
                "data_rows": row_count,
                "total_cols": len(header_clean),
                "sensitive_cols": sensitive_cols,
            })
        wb.close()
    except Exception as e:
        result["error"] = str(e)
    return result

def extract_sku_rows(filepath, sheet_name, max_rows=500):
    """提取指定工作表的SKU数据行"""
    rows_data = []
    try:
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        ws = wb[sheet_name]
        all_rows = list(ws.iter_rows(values_only=True))
        if len(all_rows) < 2:
            wb.close()
            return rows_data
        header = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(all_rows[0])]
        for row in all_rows[1:max_rows+1]:
            row_dict = {}
            for i, val in enumerate(row):
                if i < len(header):
                    col = header[i]
                    if val is not None:
                        row_dict[col] = str(val).strip()
                    else:
                        row_dict[col] = ""
            rows_data.append(row_dict)
        wb.close()
    except Exception as e:
        print(f"  Error extracting {filepath}/{sheet_name}: {e}")
    return rows_data, header

# ========== 主流程 ==========
print("=" * 60)
print("第一步：读取主xlsx文件结构")
print("=" * 60)

all_file_info = []
all_sku_records = []  # 合并所有SKU记录
sensitive_registry = []  # 敏感字段登记

for fpath in MAIN_XLSX + SUMMARY_XLSX:
    if not os.path.exists(fpath):
        print(f"  [缺失] {fpath}")
        continue
    print(f"\n处理: {os.path.basename(fpath)} ({round(os.path.getsize(fpath)/1024/1024,2)}MB)")
    info = read_xlsx_info(fpath)
    all_file_info.append(info)
    for sh in info["sheets"]:
        print(f"  Sheet: {sh['sheet_name']} | 行数: {sh['data_rows']} | 列数: {sh['total_cols']}")
        print(f"  表头: {sh['header'][:15]}")
        if sh["sensitive_cols"]:
            print(f"  ⚠️ 敏感列: {sh['sensitive_cols']}")
            sensitive_registry.append({
                "file": os.path.basename(fpath),
                "sheet": sh["sheet_name"],
                "sensitive_columns": sh["sensitive_cols"]
            })
        
        # 提取SKU行数据
        rows, header = extract_sku_rows(fpath, sh["sheet_name"])
        for r in rows:
            r["_source_file"] = os.path.basename(fpath)
            r["_source_sheet"] = sh["sheet_name"]
            all_sku_records.append(r)

print(f"\n总SKU记录数（跨所有文件）: {len(all_sku_records)}")
print(f"敏感字段登记条目: {len(sensitive_registry)}")

# 保存中间结果
with open(os.path.join(OUT_DIR, "_intermediate_sku.json"), "w", encoding="utf-8") as f:
    json.dump({
        "file_info": all_file_info,
        "total_records": len(all_sku_records),
        "sensitive_registry": sensitive_registry,
    }, f, ensure_ascii=False, indent=2)

# 保存SKU记录（不含敏感值）
safe_records = []
for r in all_sku_records:
    safe = {}
    for k, v in r.items():
        if k.startswith("_"):
            safe[k] = v
        elif is_sensitive_col(k):
            safe[k] = "[CONFIDENTIAL-存在]"
        else:
            safe[k] = v
    safe_records.append(safe)

with open(os.path.join(OUT_DIR, "_intermediate_sku_records.json"), "w", encoding="utf-8") as f:
    json.dump(safe_records, f, ensure_ascii=False, indent=2)

print("\n中间结果已保存。")
